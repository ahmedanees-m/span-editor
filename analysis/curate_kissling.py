"""Build corpus rows from the Kissling supplementary tables."""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path

SHEET = re.compile(r"^(?P<context>.+)_(?P<cas>SpCas9|SpG|SpRY)-(?P<deaminase>ABE8e|ABEmax)$")

EFFECTOR = {"ABE8e": "tada8e", "ABEmax": "abemax"}
LABEL = {"SpCas9": "none", "SpG": "scaffold_perturbed", "SpRY": "scaffold_perturbed"}
STRUCTURE = {"SpCas9": "5F9R", "SpG": "5F9R", "SpRY": "5F9R"}
LINKER_RESIDUES = 32
MARKER = "noA"


def positional_means(worksheet, spacer_length: int) -> tuple[list[float], list[int]]:
    """Mean editing efficiency at each protospacer position, over targets."""
    header = [str(v) if v is not None else "" for v in next(worksheet.iter_rows(
        min_row=1, max_row=1, values_only=True))]
    columns = {}
    for index, name in enumerate(header):
        match = re.match(r"^A(\d+)_efficiency$", name)
        if match:
            columns[int(match.group(1))] = index

    totals = {position: 0.0 for position in columns}
    counts = {position: 0 for position in columns}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        for position, index in columns.items():
            value = row[index]
            if value is None or value == MARKER:
                continue
            try:
                totals[position] += float(value)
            except (TypeError, ValueError):
                continue
            counts[position] += 1

    means, sizes = [], []
    for position in range(1, spacer_length + 1):
        count = counts.get(position, 0)
        means.append(totals[position] / count if count else 0.0)
        sizes.append(count)
    return means, sizes


def format_profile(means: list[float], sizes: list[int], minimum: int) -> str:
    return ";".join(
        f"{index}:{value:.4g}"
        for index, (value, count) in enumerate(zip(means, sizes), start=1)
        if count >= minimum
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--spacer-length", type=int, default=20)
    parser.add_argument(
        "--min-targets",
        type=int,
        default=100,
        help="positions supported by fewer targets than this are left out",
    )
    parser.add_argument("--source-key", default="kissling2025")
    parser.add_argument("--out", type=Path, default=Path("data/architectures.tsv"))
    parser.add_argument("--summary", type=Path, default=Path("results/kissling_curation.json"))
    args = parser.parse_args(argv)

    import openpyxl

    book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    rows, summary = [], []
    for name in book.sheetnames:
        match = SHEET.match(name)
        if not match:
            continue
        context = match.group("context")
        cas = match.group("cas")
        deaminase = match.group("deaminase")

        worksheet = book[name]
        means, sizes = positional_means(worksheet, args.spacer_length)
        profile = format_profile(means, sizes, args.min_targets)
        if not profile:
            continue

        peak = max(range(len(means)), key=lambda i: means[i]) + 1
        rows.append(
            {
                "entry_id": f"kissling-{cas}-{deaminase}-{context}".replace("_", "-"),
                "source_key": args.source_key,
                "editor": f"{cas}-{deaminase}",
                "ortholog": cas,
                "anchor_site": "n_terminal",
                "linker_residues": LINKER_RESIDUES,
                "linker_composition": "xten",
                "effector": EFFECTOR[deaminase],
                "label": "effector_varied" if deaminase == "ABE8e" else LABEL[cas],
                "readout_class": "amplicon_sequencing",
                "tier": "A",
                "spacer_length": args.spacer_length,
                "numbering": "pam_distal",
                "structure_id": STRUCTURE[cas],
                "split": "fit" if (cas == "SpCas9" and deaminase == "ABE8e") else "held_out",
                "context": context,
                "panel": f"Additional file 4, sheet {name}",
                "observed_profile": profile,
            }
        )
        summary.append(
            {
                "sheet": name,
                "cas": cas,
                "deaminase": deaminase,
                "context": context,
                "targets": worksheet.max_row - 1,
                "positions_kept": profile.count(";") + 1,
                "peak_position": peak,
                "peak_efficiency": round(means[peak - 1], 3),
                "targets_per_position": sizes,
            }
        )
    book.close()

    if not rows:
        print("no per-position sheets matched")
        return 1

    # Rows from this source replace any already in the corpus, so re-running the
    # curation is idempotent and leaves entries from other sources alone.
    fieldnames = list(rows[0])
    existing = []
    if args.out.exists():
        with open(args.out, newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            if reader.fieldnames:
                fieldnames = list(dict.fromkeys(list(reader.fieldnames) + fieldnames))
            existing = [row for row in reader if row.get("source_key") != args.source_key]

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with open(args.out, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=fieldnames, delimiter="\t", extrasaction="ignore"
        )
        writer.writeheader()
        writer.writerows(existing + rows)

    args.summary.parent.mkdir(parents=True, exist_ok=True)
    with open(args.summary, "w", encoding="utf-8") as handle:
        json.dump({"workbook": str(args.workbook), "entries": summary}, handle, indent=2)
        handle.write("\n")

    print(f"{len(rows)} entries written to {args.out}")
    print()
    print("architecture           context                  targets  peak  efficiency")
    for record in summary:
        print(
            f"{record['cas']:8s} {record['deaminase']:8s} {record['context']:24s} "
            f"{record['targets']:7d}  {record['peak_position']:4d}  "
            f"{record['peak_efficiency']:9.3f}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
